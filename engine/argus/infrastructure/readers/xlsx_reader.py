from collections.abc import Iterator
from datetime import date, datetime, time
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook

from argus.application.ports.document_reader import DocumentSource
from argus.domain.document import Document, DocumentFormat
from argus.domain.errors import UnreadableDocumentError, UnreadableReason
from argus.domain.locations import SheetCellLocation, SourceLocation
from argus.infrastructure.readers.fragments import build_fragments, require_text_found

MIDNIGHT = time()


class XlsxDocumentReader:
    """Lee el valor calculado de cada celda con datos; el origen es hoja y coordenada (C12)."""

    def read(self, source: DocumentSource) -> Document:
        workbook = _open_workbook(source)
        try:
            fragments = build_fragments(_cell_texts(workbook))
        finally:
            workbook.close()
        return Document(
            source.name, DocumentFormat.XLSX, require_text_found(source.name, fragments)
        )


def _open_workbook(source: DocumentSource) -> Workbook:
    try:
        # data_only: el valor que ve el usuario, no la fórmula que lo produce.
        return load_workbook(BytesIO(source.content), read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, KeyError, OSError):
        raise UnreadableDocumentError(source.name, UnreadableReason.CORRUPTED) from None


def _cell_texts(workbook: Workbook) -> Iterator[tuple[str, SourceLocation]]:
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            yield from (
                (format_cell_value(cell.value), SheetCellLocation(sheet.title, cell.coordinate))
                for cell in row
                if cell.value is not None
            )


def format_cell_value(value: object) -> str:
    """Excel guarda las fechas como fecha y hora; si la hora es 00:00 se muestra solo la fecha."""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == MIDNIGHT else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)
