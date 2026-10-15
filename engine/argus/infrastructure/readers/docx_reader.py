from collections.abc import Iterator
from io import BytesIO
from itertools import chain
from zipfile import BadZipFile

import docx
from docx.document import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError

from argus.application.ports.document_reader import DocumentSource
from argus.domain.document import Document, DocumentFormat
from argus.domain.errors import UnreadableDocumentError, UnreadableReason
from argus.domain.locations import ParagraphLocation, SourceLocation, TableCellLocation
from argus.infrastructure.readers.fragments import build_fragments, require_text_found

FIRST_POSITION = 1


class DocxDocumentReader:
    """Lee párrafos y celdas de tabla con python-docx.

    Un DOCX no guarda páginas (las calcula Word al mostrarlo), así que el origen se indica por
    párrafo o por tabla, fila y columna.
    """

    def read(self, source: DocumentSource) -> Document:
        word = _open_word_document(source)
        texts = chain(_paragraph_texts(word), _table_cell_texts(word))
        fragments = require_text_found(source.name, build_fragments(texts))
        return Document(source.name, DocumentFormat.DOCX, fragments)


def _open_word_document(source: DocumentSource) -> WordDocument:
    try:
        return docx.Document(BytesIO(source.content))
    except (PackageNotFoundError, BadZipFile, KeyError, ValueError):
        raise UnreadableDocumentError(source.name, UnreadableReason.CORRUPTED) from None


def _paragraph_texts(word: WordDocument) -> Iterator[tuple[str, SourceLocation]]:
    for number, paragraph in enumerate(word.paragraphs, start=FIRST_POSITION):
        yield paragraph.text, ParagraphLocation(number)


def _table_cell_texts(word: WordDocument) -> Iterator[tuple[str, SourceLocation]]:
    for table_number, table in enumerate(word.tables, start=FIRST_POSITION):
        for row_number, row in enumerate(table.rows, start=FIRST_POSITION):
            yield from (
                (cell.text, TableCellLocation(table_number, row_number, column_number))
                for column_number, cell in enumerate(row.cells, start=FIRST_POSITION)
            )
