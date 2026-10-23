from io import BytesIO

from docx import Document as WordDocument
from docx.document import Document as WordFile
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from argus.domain.report_document import (
    PAGE_PLACEHOLDER,
    TOTAL_PLACEHOLDER,
    ReportCell,
    ReportDocument,
    ReportSection,
    ReportTable,
)

DETAIL_COLOR = RGBColor(0x4F, 0x4B, 0x44)
DETAIL_SIZE = Pt(8)
TABLE_STYLE = "Table Grid"
FIELDS = {PAGE_PLACEHOLDER: "PAGE", TOTAL_PLACEHOLDER: "NUMPAGES"}


def _field(paragraph: Paragraph, instruction: str) -> None:
    """Un campo de Word que se actualiza solo: el número de página o el total."""
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instruction), ("end", None)):
        element = OxmlElement("w:fldChar" if kind else "w:instrText")
        if kind:
            element.set(qn("w:fldCharType"), kind)
        else:
            element.text = text
        run._r.append(element)


def _page_label(paragraph: Paragraph, label: str) -> None:
    """«Página {page} de {total}» con los números como campos de Word."""
    text = label
    while text:
        positions = [(text.find(token), token) for token in FIELDS if token in text]
        if not positions:
            paragraph.add_run(text)
            return
        index, token = min(positions)
        paragraph.add_run(text[:index])
        _field(paragraph, FIELDS[token])
        text = text[index + len(token) :]


def _detail(paragraph: Paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    run.font.size = DETAIL_SIZE
    run.font.color.rgb = DETAIL_COLOR


def _cell(cell: _Cell, value: ReportCell) -> None:
    cell.paragraphs[0].add_run(value.text)
    if value.detail:
        _detail(cell.add_paragraph(), value.detail)


def _table(word: WordFile, table: ReportTable) -> None:
    grid = word.add_table(rows=1, cols=len(table.columns))
    grid.style = TABLE_STYLE
    for cell, column in zip(grid.rows[0].cells, table.columns, strict=True):
        cell.paragraphs[0].add_run(column).bold = True
    for row in table.rows:
        for cell, value in zip(grid.add_row().cells, row, strict=True):
            _cell(cell, value)


def _section(word: WordFile, section: ReportSection) -> None:
    word.add_heading(section.heading, level=2)
    if section.metrics:
        line = word.add_paragraph()
        for metric in section.metrics:
            line.add_run(f"{metric.value} ").bold = True
            line.add_run(f"{metric.label}    ")
    for text in section.paragraphs:
        word.add_paragraph(text)
    if section.table is not None:
        _table(word, section.table)


class DocxReportRenderer:
    """Maqueta el informe en DOCX con los estilos de Word, para que se pueda editar después."""

    def render(self, document: ReportDocument) -> bytes:
        word = WordDocument()
        word.core_properties.title = document.title
        word.add_heading(document.title, level=1)
        for line in (*document.subtitles, document.brand):
            _detail(word.add_paragraph(), line)
        for section in document.sections:
            _section(word, section)
        footer = word.sections[0].footer.paragraphs[0]
        footer.add_run(f"{document.footer}    ")
        _page_label(footer, document.page_label)
        output = BytesIO()
        word.save(output)
        return output.getvalue()
