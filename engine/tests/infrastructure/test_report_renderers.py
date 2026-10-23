from io import BytesIO
from pathlib import Path

from docx import Document as WordDocument
from pypdf import PdfReader

from argus.infrastructure.reports.documents_report_folder import DocumentsReportFolder
from argus.infrastructure.reports.docx_report_renderer import DocxReportRenderer
from argus.infrastructure.reports.pdf_report_renderer import PdfReportRenderer
from tests.domain.report_builders import report_document


def test_the_pdf_has_every_section_its_evidence_and_the_page_footer() -> None:
    content = PdfReportRenderer().render(report_document())

    text = "".join(page.extract_text() for page in PdfReader(BytesIO(content)).pages)
    assert content.startswith(b"%PDF")
    for expected in (
        "Informe de revisión documental",
        "1. Resumen",
        "Contrato_4857_v3.docx · Página 1",
        "Generado con ARGUS 0.1.0",
        "Página 1 de 1",
    ):
        assert expected in text


def test_the_docx_can_be_edited_and_numbers_its_pages() -> None:
    content = DocxReportRenderer().render(report_document())

    word = WordDocument(BytesIO(content))
    headings = [
        p.text
        for p in word.paragraphs
        if p.style is not None and p.style.name.startswith("Heading")
    ]
    assert headings == ["Informe de revisión documental", "1. Resumen", "2. Evidencia"]
    (table,) = word.tables
    assert table.rows[1].cells[1].text == "4857\nContrato_4857_v3.docx · Página 1"
    footer = word.sections[0].footer.paragraphs[0]
    assert "Generado con ARGUS" in footer.text
    assert "w:instrText" in footer._p.xml
    assert "NUMPAGES" in footer._p.xml


def test_the_folder_keeps_each_case_file_apart_and_never_leaves_it(tmp_path: Path) -> None:
    folder = DocumentsReportFolder(tmp_path)

    saved = folder.save("Expediente_2026-118", "../Informe.pdf", b"%PDF")

    assert saved.folder == str(tmp_path / "Expediente_2026-118")
    assert (tmp_path / "Expediente_2026-118" / "Informe.pdf").read_bytes() == b"%PDF"
    assert folder.path_of("Expediente_2026-118", "Informe.pdf") is not None
    assert folder.path_of("Expediente_2026-118", "otro.pdf") is None
