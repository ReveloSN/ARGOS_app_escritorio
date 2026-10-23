"""Genera los archivos de ejemplo de tests/fixtures con los datos del expediente de los mockups.

Se ejecuta a mano cuando cambian los fixtures y el resultado se versiona:
    uv run python -m tests.fixtures.generate_fixtures
"""

from datetime import date
from email.message import EmailMessage
from io import BytesIO
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document as WordDocument
from openpyxl import Workbook
from pypdf import PdfWriter

FIXTURES_DIR = Path(__file__).resolve().parent
PDF_USER_PASSWORD = "secreto"
PDF_OWNER_PASSWORD = "propietario"
LETTER_WIDTH, LETTER_HEIGHT = 612, 792

CONTRACT_PAGES = [
    [
        "CONTRATO DE ARRENDAMIENTO DE BODEGA COMERCIAL",
        "Celebran el presente Contrato No. 4857 de arrendamiento de bodega comercial,",
        "ubicada en la Calle 17 No. 96G-45, Fontibón, Bogotá D.C.",
        "Arrendador: Inversiones Torres Gómez S.A.S., NIT 900.482.315-7",
    ],
    [
        "CLÁUSULA TERCERA. Canon mensual: $ 18.450.000",
        "Fecha de inicio: 01/07/2026. Vigencia: 36 meses.",
        "Representante legal: María F. Rojas Pérez",
    ],
]


def escape_pdf_text(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


SCANNED_ANNEX_LINES = [
    "ANEXO 4. ACTA DE ENTREGA DE LA BODEGA",
    "Hace parte del Contrato No. 4587 de arrendamiento.",
    "Canon mensual: $ 18.450.000",
]
# Resolución de un escaneo de oficina: suficiente para el OCR y un archivo liviano.
SCAN_DPI = 200
POINTS_PER_INCH = 72


def build_content_stream(lines: list[str]) -> bytes:
    commands = ["BT", "/F1 11 Tf", "72 720 Td", "16 TL"]
    commands += [f"({escape_pdf_text(line)}) Tj T*" for line in lines]
    commands.append("ET")
    return "\n".join(commands).encode("cp1252")


def build_pdf_objects(pages: list[list[str]]) -> dict[int, bytes]:
    page_ids = [4 + 2 * index for index in range(len(pages))]
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode(),
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    }
    for page_id, lines in zip(page_ids, pages, strict=True):
        stream = build_content_stream(lines)
        objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {LETTER_WIDTH} {LETTER_HEIGHT}] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {page_id + 1} 0 R >>"
        ).encode()
        objects[page_id + 1] = b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream)
    return objects


def build_text_pdf(pages: list[list[str]]) -> bytes:
    """PDF mínimo con texto real en Helvetica, escrito a mano para no depender de otra librería."""
    objects = build_pdf_objects(pages)
    output = bytearray(b"%PDF-1.4\n")
    offsets = []
    for object_id in sorted(objects):
        offsets.append(len(output))
        output += b"%d 0 obj\n%s\nendobj\n" % (object_id, objects[object_id])
    xref_offset = len(output)
    output += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    output += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    output += b"trailer\n<< /Size %d /Root 1 0 R >>\n" % (len(objects) + 1)
    output += b"startxref\n%d\n%%%%EOF\n" % xref_offset
    return bytes(output)


def write_pdf_fixtures() -> None:
    contract = build_text_pdf(CONTRACT_PAGES)
    (FIXTURES_DIR / "Contrato_4857_v3.pdf").write_bytes(contract)

    protected = PdfWriter(clone_from=BytesIO(contract))
    protected.encrypt(PDF_USER_PASSWORD, PDF_OWNER_PASSWORD, algorithm="RC4-128")
    protected.write(FIXTURES_DIR / "Estados_financieros_Q2.pdf")

    scanned = PdfWriter()
    scanned.add_blank_page(LETTER_WIDTH, LETTER_HEIGHT)
    scanned.write(FIXTURES_DIR / "anexo_3.pdf")

    (FIXTURES_DIR / "danado.pdf").write_bytes(b"%PDF-1.4\nesto no es un PDF completo")


def build_scanned_pdf(lines: list[str]) -> bytes:
    """PDF solo con la imagen de la página, sin capa de texto: como lo entrega un escáner."""
    page = pdfium.PdfDocument(build_text_pdf([lines]))[0]
    image = page.render(scale=SCAN_DPI / POINTS_PER_INCH).to_pil().convert("L")
    output = BytesIO()
    image.save(output, format="PDF", resolution=SCAN_DPI)
    return output.getvalue()


def write_scanned_fixture() -> None:
    scanned = build_scanned_pdf(SCANNED_ANNEX_LINES)
    (FIXTURES_DIR / "Anexo_4_escaneado.pdf").write_bytes(scanned)


def build_annex_docx() -> bytes:
    document = WordDocument()
    document.add_paragraph("ANEXO 1 · CONDICIONES GENERALES")
    document.add_paragraph("")
    document.add_paragraph(
        "El presente anexo hace parte integral del Contrato No. 4587 suscrito entre "
        "Inversiones Torres Gómez S.A.S. y Logística Andina S.A.S., NIT 900.482.315-7."
    )
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Concepto"
    table.cell(0, 1).text = "Valor"
    table.cell(1, 0).text = "Canon mensual"
    table.cell(1, 1).text = "$ 18.540.000"
    document.add_paragraph("Representante legal: María F. Rojas Peña")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_rates_xlsx() -> bytes:
    workbook = Workbook()
    summary = workbook.active
    assert summary is not None
    summary.title = "Resumen"
    summary["A1"] = "Tarifas mayo 2026"
    summary["B12"] = "Canon mensual"
    summary["C12"] = 18540000
    summary["C14"] = date(2026, 5, 31)
    workbook.create_sheet("Notas")["A1"] = "Valores sin IVA"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


NOTES_TEXT = "Notas de la remisión\n\nInforme de gestión de junio 2026\nCanon: $ 18.450.000\n"


def build_remission_email(rates: bytes) -> bytes:
    message = EmailMessage()
    message["From"] = "laura.martinez@example.com"
    message["To"] = "cliente@example.com"
    message["Subject"] = "Remisión del informe de junio"
    message.set_content("Buenos días:\n\nAdjunto el informe de junio con las tarifas.\n")
    message.add_attachment(
        rates,
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="Anexo_2_tarifas.xlsx",
    )
    message.add_attachment(NOTES_TEXT, filename="notas.txt")
    message.add_attachment(
        b"binario", maintype="application", subtype="octet-stream", filename="Acta_comite_oct.pages"
    )
    return message.as_bytes()


# Verificaciones internas (ADR 022): cifra en letras distinta, fecha imposible y campo vacío.
DRAFT_TEXT = (
    "Borrador del Contrato No. 4857\n"
    "Canon mensual de $ 18.450.000 (dieciocho millones quinientos cuarenta mil pesos m/cte).\n"
    "Fecha de entrega: 31/02/2026.\n"
    "Depósito en garantía: $ ________\n"
)


def write_office_and_text_fixtures() -> None:
    rates = build_rates_xlsx()
    (FIXTURES_DIR / "Anexo_1_condiciones.docx").write_bytes(build_annex_docx())
    (FIXTURES_DIR / "Anexo_2_tarifas.xlsx").write_bytes(rates)
    (FIXTURES_DIR / "Correo_remision_cliente.eml").write_bytes(build_remission_email(rates))
    (FIXTURES_DIR / "notas_utf8.txt").write_bytes(NOTES_TEXT.encode("utf-8"))
    (FIXTURES_DIR / "notas_windows.txt").write_bytes(NOTES_TEXT.encode("cp1252"))
    (FIXTURES_DIR / "Borrador_contrato.txt").write_bytes(DRAFT_TEXT.encode("utf-8"))
    (FIXTURES_DIR / "danado.docx").write_bytes(b"PK no es un docx")
    (FIXTURES_DIR / "danado.xlsx").write_bytes(b"PK no es un xlsx")


if __name__ == "__main__":
    write_pdf_fixtures()
    write_scanned_fixture()
    write_office_and_text_fixtures()
